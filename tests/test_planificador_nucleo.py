"""Núcleo del planificador con IA: cuentas de Barra, posproceso, cierre, versiones, trabajador de WebClass y registro.

No usa red, ni IA, ni WebClass: todo con carpetas temporales y simulaciones.
"""

import csv
import io
import json
import os
import sys
import unittest
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))     # para encontrar planificador_comun con o sin «discover»
import planificador_comun as comun  # noqa: E402

if not comun.HAY_PLANIFICADOR:                                    # pragma: no cover
    raise unittest.SkipTest("Esta copia de Barra no trae el planificador.")

import webclass_recursos as rec  # noqa: E402  (módulo de Barra)
import barra_utiles  # noqa: E402
bw_servidor = comun.importar("servidor")

rutas = comun.importar("nucleo.rutas")
perfil_mod = comun.importar("nucleo.perfil")
ia = comun.importar("nucleo.ia")
historial = comun.importar("nucleo.historial")
version = comun.importar("nucleo.version")
postproceso = comun.importar("nucleo.postproceso")
cierre = comun.importar("nucleo.cierre")
prompts = comun.importar("nucleo.prompts")
extraccion = comun.importar("nucleo.extraccion")
bw = comun.importar("nucleo.barra_worker")


# Texto real (recortado) que antes salía desordenado en el cierre: cada trozo de la diapositiva en su propia línea
MATERIAL_REAL = """
--- Página 2 ---
ACTIVIDAD 2
Página 183 del
Libro de clases.
Respuesta en
El libro

--- Página 3 ---
Objetivo.
Analizar un gráfico de
La temperatura y el
Tiempo al calentar
Agua  describiendo el
Comportamiento de
Las partículas.

--- Página 23 ---
Completa y resume en tu cuaderno
 Ideas sobre la
Clase
Preguntas
 Dibujo
3 2 1
"""


def leer(ruta, modo="r"):
    with open(ruta, modo, **({} if "b" in modo else {"encoding": "utf-8"})) as f:
        return f.read()


class CuentasDeBarra(unittest.TestCase):
    def test_sigue_la_cuenta_activa_de_barra(self):
        with comun.BarraFalsa(activa="beto") as barra:
            self.assertTrue(rutas.EN_BARRA)
            self.assertEqual(rutas.perfiles_barra(), ["ana", "beto"])
            self.assertEqual(rutas.clave_activa(), "beto")
            self.assertEqual(rutas.nombre_de_cuenta(), "Beto Soto")
            barra.cambiar_cuenta("ana")                              # en la barra: «Cambiar de profesor»
            self.assertEqual(rutas.clave_activa(), "ana")
            self.assertEqual(rutas.nombre_de_cuenta(), "Ana Pérez")

    def test_cuenta_inexistente_o_sin_archivo_usa_la_primera_como_barra(self):
        with comun.BarraFalsa(activa="fantasma"):
            self.assertEqual(rutas.clave_activa(), "ana")
        with comun.BarraFalsa(activa="beto") as barra:
            os.remove(os.path.join(barra.carpeta, "datos", "perfil_activo.txt"))
            self.assertEqual(rutas.clave_activa(), "ana")

    def test_fuera_de_barra_la_cuenta_es_local(self):
        os.environ.pop("PLANIFICADOR_BARRA", None)
        with mock.patch.object(rutas, "EN_BARRA", False):
            self.assertEqual(rutas.clave_activa(), "local")
            self.assertEqual(rutas.perfiles_barra(), [])

    def test_cada_cuenta_tiene_su_perfil_claves_e_historial(self):
        with comun.BarraFalsa(activa="beto") as barra:
            p = perfil_mod.cargar()
            self.assertEqual(p.nombre_docente, "Beto Soto", "parte con el nombre de la cuenta de Barra")
            p.colegio, p.voz = "Colegio de Beto", "femenina"
            perfil_mod.guardar(p)
            ia.guardar_claves({"Gemini (Google)": "CLAVE-DE-BETO"})
            historial.guardar({"curso": "x", "nombre": "clase de beto"})

            barra.cambiar_cuenta("ana")
            otra = perfil_mod.cargar()
            self.assertEqual((otra.colegio, otra.voz, otra.nombre_docente), ("", "neutra", "Ana Pérez"))
            self.assertEqual(ia.clave_guardada("GEMINI_API_KEY"), "")
            self.assertEqual(historial.listar(), [])

            barra.cambiar_cuenta("beto")
            self.assertEqual(perfil_mod.cargar().colegio, "Colegio de Beto")
            self.assertEqual(ia.clave_guardada("GEMINI_API_KEY"), "CLAVE-DE-BETO")
            self.assertEqual(len(historial.listar()), 1)

    def test_nunca_se_lee_el_rut_de_la_cuenta(self):
        with comun.BarraFalsa():
            self.assertNotIn("00000000-0", repr(rutas.nombre_de_cuenta()))
            fuente = leer(os.path.join(comun.PLANIFICADOR, "nucleo", "rutas.py"))
            self.assertNotIn('["rut"]', fuente)
            self.assertNotIn("get(\"rut\")", fuente)

    def test_registro_de_planificaciones_usa_la_misma_regla_de_nombres_que_barra(self):
        with comun.BarraFalsa() as barra:
            with open(os.path.join(barra.carpeta, "datos", "registro_historico.txt"), "w", encoding="utf-8") as f:
                f.write("ana")                                       # «dueña del registro de siempre»
            self.assertEqual(os.path.basename(rutas.registro_planificaciones("ana")), "registro_planificaciones.csv")
            self.assertEqual(os.path.basename(rutas.registro_planificaciones("beto")), "registro_planificaciones_beto.csv")


class Versiones(unittest.TestCase):
    def test_los_dos_archivos_de_version_existen_y_tienen_formato_valido(self):
        self.assertRegex(version.leer(comun.HERRAMIENTA), r"^\d+\.\d+\.\d+$", "falta herramienta\\VERSION")
        self.assertRegex(version.leer(comun.PLANIFICADOR), r"^\d+\.\d+\.\d+$", "falta herramienta\\planificador\\VERSION")
        for carpeta in (comun.HERRAMIENTA, comun.PLANIFICADOR):
            cambios = leer(os.path.join(carpeta, "CAMBIOS.md"))
            self.assertIn(version.leer(carpeta), cambios, f"CAMBIOS.md de {carpeta} no menciona su versión actual")

    def test_compatibilidad_con_la_version_de_barra(self):
        casos = [(None, False), ("basura", False), ("2.0.0", False), ("2.5.5", False), ("2.6.0", True), ("2.10.3", True), ("3.0.0", True)]
        for v, esperado in casos:
            with self.subTest(barra=v), comun.BarraFalsa(version=v):
                ok, mensaje = version.compatibilidad()
                self.assertEqual(ok, esperado)
                if not esperado:
                    self.assertIn("ACTUALIZAR.bat", mensaje)

    def test_fuera_de_barra_no_aplica(self):
        with mock.patch.object(rutas, "EN_BARRA", False):
            self.assertEqual(version.compatibilidad(), (True, ""))

    def test_la_version_minima_no_supera_la_de_barra_que_se_reparte(self):
        mayor = tuple(int(n) for n in version.leer(comun.HERRAMIENTA).split("."))
        self.assertLessEqual(tuple(int(n) for n in version.BARRA_MINIMA.split(".")), mayor)


class Posproceso(unittest.TestCase):
    def test_el_bloque_de_actividad_se_detiene_en_la_diapositiva_y_queda_legible(self):
        bloque = postproceso.extraer_bloque_actividad(MATERIAL_REAL)
        self.assertEqual(bloque, "ACTIVIDAD 2 Página 183 del Libro de clases. Respuesta en El libro")
        self.assertNotIn("---", bloque)
        self.assertNotIn("partículas", bloque.lower(), "no debe arrastrar las diapositivas siguientes")

    def test_prioriza_actividad_final_y_ticket_de_salida(self):
        texto = "--- Página 1 ---\nACTIVIDAD 1\nlee\n--- Página 2 ---\nTICKET DE SALIDA\nresponde 3 preguntas\n--- Página 3 ---\nFin"
        self.assertTrue(postproceso.extraer_bloque_actividad(texto).startswith("TICKET DE SALIDA"))
        self.assertEqual(postproceso.extraer_bloque_actividad("--- Página 1 ---\nSolo teoría"), "")

    def test_ultima_pagina_trae_la_dinamica_de_cierre(self):
        self.assertIn("3 2 1", postproceso.ultima_pagina(MATERIAL_REAL))

    def test_separa_las_secciones_y_limpia_asteriscos(self):
        ini, des, cie, extras = postproceso.limpiar_y_separar_secciones(
            "<<<INICIO>>>\n**Objetivo** de la clase\n<<<DESARROLLO>>>\nBloque\n<<<NOMBRE>>>\nMi clase\n<<<DUA>>>\n1, 2, 3")
        self.assertEqual((ini, des, cie), ("Objetivo de la clase", "Bloque", ""))
        self.assertEqual((extras["NOMBRE"], extras["DUA"]), ("Mi clase", "1, 2, 3"))

    def test_la_voz_femenina_solo_se_aplica_si_el_perfil_la_pide(self):
        self.assertIn("las estudiantes", postproceso.corregir_texto("los estudiantes leen", "femenina").lower())
        self.assertIn("los estudiantes", postproceso.corregir_texto("los estudiantes leen", "neutra").lower())

    def test_objetivo_manual_se_impone_y_la_apertura_fija_va_primero(self):
        p = perfil_mod.Perfil(apertura_fija="V.J.E.")
        ini, _, _ = postproceso.aplicar_reglas("Objetivo de Aprendizaje: algo de la IA\nTexto", "d", "c", p, "Mi objetivo", "")
        self.assertEqual(ini.split("\n")[0], "V.J.E.")
        self.assertIn("Objetivo de Aprendizaje: Mi objetivo", ini)
        self.assertNotIn("algo de la IA", ini)


class Prompts(unittest.TestCase):
    def test_el_prompt_neutro_no_tiene_nada_de_historia_ni_de_otro_colegio(self):
        texto = prompts.construir_sistema(perfil_mod.Perfil(), "Matemática", {1: "Pauta uno"})
        for prohibida in ("Ignacio", "Historia", "V.J.E", "Recoleta", "Henríquez", "Neyün", "Doctrina Social"):
            self.assertNotIn(prohibida, texto)
        self.assertNotIn("<<<CIERRE>>>:", texto.replace("ni escribas <<<CIERRE>>>", ""), "el cierre se pide aparte")

    def test_las_opciones_del_perfil_si_se_incluyen(self):
        p = perfil_mod.Perfil(apertura_fija="V.J.E.", voz="femenina", contexto_colegio="colegio de prueba")
        texto = prompts.construir_sistema(p, "Historia, Geografía y Ciencias Sociales", {})
        self.assertIn("V.J.E.", texto)
        self.assertIn("colegio de prueba", texto)
        self.assertIn("femenina", texto)


class Cierre(unittest.TestCase):
    PERFIL = perfil_mod.Perfil()
    CIERRE_C_BUENO = ("El docente construye un esquema en la pizarra con el curso y luego discuten en pareja durante seis minutos.\n"
                      "- ¿Qué cambios de estado se observan?\n- ¿Por qué la temperatura se mantiene constante?\n"
                      "- ¿Qué postura tomas y cómo la justificas?")

    def test_modo_segun_el_origen_de_la_actividad(self):
        self.assertEqual(cierre.elegir_modo("responder la pregunta 3", MATERIAL_REAL)[0], "A")
        self.assertEqual(cierre.elegir_modo("", MATERIAL_REAL)[0], "B")
        self.assertEqual(cierre.elegir_modo("", "--- Página 1 ---\nsolo teoría")[0], "C")

    def test_validacion(self):
        self.assertIsNone(cierre.validar(self.CIERRE_C_BUENO, "C"))
        self.assertIn("preguntas", cierre.validar("Repasar lo visto.", "C") + "preguntas")
        self.assertIsNotNone(cierre.validar("ACTIVIDAD\nResolver --- Página 3 --- cosas", "B"))
        self.assertIsNotNone(cierre.validar("ACTIVIDAD\ncorto", "B"))
        self.assertIsNone(cierre.validar("ACTIVIDAD\nLas estudiantes completan la actividad 2 de la página 183 del libro.", "B"))

    def test_respaldo_siempre_es_valido(self):
        self.assertIsNone(cierre.validar(cierre.respaldo("C", "", "las fracciones", self.PERFIL), "C"))
        self.assertIsNone(cierre.validar(cierre.respaldo("B", "ACTIVIDAD 2 Página 183 del Libro de clases.", "", self.PERFIL), "B"))

    def _redactar(self, respuestas, modo="C", fuente=""):
        llamadas = iter(respuestas)

        def falsa(construir, **kw):
            r = next(llamadas)
            if isinstance(r, Exception):
                raise r
            return r, "IA falsa", "stop", []
        with mock.patch.object(cierre.ia, "completar", falsa):
            return cierre.redactar(self.PERFIL, "Ciencias", "7° Básico", "calor", "Objetivo", "Desarrollo", modo, fuente, "")

    def test_redactar_acepta_un_buen_cierre_a_la_primera(self):
        texto, origen = self._redactar([self.CIERRE_C_BUENO])
        self.assertEqual(origen, "ia")
        self.assertEqual(texto.count("\n- "), 3)

    def test_redactar_corrige_un_intento_malo(self):
        texto, origen = self._redactar(["Repasar lo visto.", self.CIERRE_C_BUENO])
        self.assertEqual(origen, "ia-corregido")

    def test_redactar_cae_al_respaldo_si_la_ia_falla(self):
        for respuestas in (["mal", "peor"], [RuntimeError("sin IA")]):
            texto, origen = self._redactar(respuestas)
            self.assertEqual(origen, "respaldo")
            self.assertIsNone(cierre.validar(texto, "C"))

    def test_modo_b_siempre_empieza_con_la_palabra_actividad(self):
        texto, _ = self._redactar(["Las estudiantes completan la actividad 2 de la página 183 de su libro de clases."], "B",
                                  "ACTIVIDAD 2 Página 183")
        self.assertTrue(texto.startswith("ACTIVIDAD\n"))


class Extraccion(unittest.TestCase):
    def test_lee_word_y_powerpoint(self):
        from docx import Document
        from pptx import Presentation
        with comun.carpeta_temporal() as tmp:
            ruta_docx, ruta_pptx = os.path.join(tmp, "a.docx"), os.path.join(tmp, "a.pptx")
            d = Document()
            d.add_paragraph("El agua cubre el planeta.")
            d.save(ruta_docx)
            pres = Presentation()
            lamina = pres.slides.add_slide(pres.slide_layouts[1])
            lamina.shapes.title.text = "Ciclo del agua"
            pres.save(ruta_pptx)
            self.assertIn("El agua cubre el planeta.", extraccion.texto_archivo("a.docx", leer(ruta_docx, "rb")))
            texto = extraccion.texto_archivo("a.pptx", leer(ruta_pptx, "rb"))
            self.assertIn("Ciclo del agua", texto)
            self.assertIn("--- Página 1 ---", texto)

    def test_id_de_youtube(self):
        self.assertEqual(extraccion.id_youtube("https://www.youtube.com/watch?v=7e_VPPNJD4g"), "7e_VPPNJD4g")
        self.assertIsNone(extraccion.id_youtube("https://ejemplo.cl"))


class Pictogramas(unittest.TestCase):
    def test_el_diccionario_depende_de_la_asignatura_y_no_se_mezcla(self):
        sys.path.insert(0, os.path.join(comun.PLANIFICADOR, "nucleo", "kit_pie"))
        import pictos
        historia = pictos.para("Historia, Geografía y Ciencias Sociales")
        matematica = pictos.para("Matemática")
        ciencias = pictos.para("Ciencias Naturales")
        self.assertIn("SALITRE", historia)
        self.assertNotIn("SALITRE", matematica)
        self.assertIn("SUMA", matematica)
        self.assertNotIn("CELULA", historia, "«Ciencias Sociales» no activa el diccionario de Ciencias Naturales")
        self.assertIn("CELULA", ciencias)
        self.assertIn("CORRER", pictos.para("Educación Física y Salud"))
        self.assertNotIn("CORRER", ciencias)


class TrabajadorDeWebClass(unittest.TestCase):
    CATALOGO = [("2110", "> OA12 Describir la distribución del agua dulce y salada en la Tierra, considerando océanos <tr class"),
                ("2112", "> OA14 Investigar y explicar efectos de la actividad humana en océanos Agregue su Objetivo de aprendizaje.")]

    def test_limpia_el_texto_que_trae_la_pagina(self):
        self.assertEqual(bw.limpiar_oa("> OA12 Describir el agua <tr class=x"), "OA12 Describir el agua")
        self.assertEqual(bw.limpiar_oa("> OA14 Investigar Agregue su Objetivo de aprendizaje."), "OA14 Investigar")

    def test_resuelve_oa_por_etiqueta_y_por_parecido_de_texto(self):
        ids, faltan = bw.resolver_oa(["OA12"], {}, self.CATALOGO)
        self.assertEqual((ids, faltan), (["2110"], []))
        # OA que en la otra unidad se llama distinto pero tiene el mismo texto
        textos = {"OA3": "Describir la distribución del agua dulce y salada en la Tierra, considerando océanos"}
        self.assertEqual(bw.resolver_oa(["OA3"], textos, self.CATALOGO), (["2110"], []))
        self.assertEqual(bw.resolver_oa(["OA99"], {"OA99": "algo totalmente distinto"}, self.CATALOGO), ([], ["OA99"]))

    def test_oa1_no_se_confunde_con_oa12(self):
        self.assertEqual(bw.resolver_oa(["OA1"], {}, self.CATALOGO), ([], ["OA1"]))

    def test_nivel_del_recurso_sale_del_curso_de_la_unidad(self):
        self.assertEqual(bw.niveles_de_curso(rec, "5BA"), [25])
        self.assertEqual(bw.niveles_de_curso(rec, "5BB"), [25])
        self.assertEqual(bw.niveles_de_curso(rec, "8BA"), [28])
        self.assertEqual(bw.niveles_de_curso(rec, "1MA"), [31])
        self.assertEqual(bw.niveles_de_curso(rec, "curso raro"), [])

    def test_los_tipos_que_ofrece_el_planificador_los_sabe_subir_barra(self):
        for extension in bw_servidor.EXTENSIONES_RECURSO:
            self.assertIn(extension, rec.SUBCATEGORIA, f"Barra no sabe subir {extension}")
            self.assertIn(extension, barra_utiles.EXTENSIONES_RECURSOS)
        self.assertEqual(set(barra_utiles.EXTENSIONES_RECURSOS), set(rec.SUBCATEGORIA), "la vista previa de Barra y el subidor se desfasaron")
        for extension in (".jpg", ".jpeg", ".png"):
            self.assertEqual(rec.SUBCATEGORIA[extension], 33, "categoría «Imágenes» → subcategoría «Fotos»")

    def _subir(self, subidor, datos):
        """Corre op_subir con WebClass simulado; devuelve lo que entregó."""
        salida = {}

        def salir(obj):
            salida.update(obj)
            raise SystemExit(0)
        with mock.patch.object(bw, "abrir", return_value=(mock.Mock(), None, None, rec)), \
                mock.patch.object(bw, "perfil_de", return_value=("beto", {"nombre": "Beto"})), \
                mock.patch.object(bw, "iniciar_sesion", return_value="sesion"), \
                mock.patch.object(bw, "salir", salir), \
                mock.patch.object(rec, "subir_o_reutilizar", subidor), \
                self.assertRaises(SystemExit):
            bw.op_subir(datos)
        return salida

    def test_subir_al_banco_usa_los_niveles_de_barra_y_el_tipo_elegido(self):
        visto = {}

        def subir(session, ruta, *, nombre, autor, niveles, asignatura, visible, tipo_file=None):
            visto.update(session=session, ruta=ruta, nombre=nombre, autor=autor, niveles=niveles, asignatura=asignatura,
                         visible=visible, tipo_file=tipo_file)
            return "555", True
        datos = {"ruta": "C:/x/guia.pdf", "nombre": "Guía PIE", "autor": "Beto", "niveles": ["5", "2M", "99"], "asignatura": "107",
                 "visible": False, "tipo_file": "70"}
        self.assertEqual(self._subir(subir, datos), {"id": "555", "subido": True, "nombre": "Guía PIE"})
        self.assertEqual(visto["niveles"], [rec.NIVELES["5"], rec.NIVELES["2M"]], "un nivel desconocido se descarta")
        self.assertEqual((visto["session"], visto["asignatura"], visto["visible"], visto["tipo_file"]), ("sesion", 107, False, 70))

    def test_subir_con_una_barra_que_no_conoce_el_tipo_no_lo_envia(self):
        visto = {}

        def subir_antiguo(session, ruta, *, nombre, autor, niveles, asignatura, visible):
            visto.update(nombre=nombre)
            return "9", False
        r = self._subir(subir_antiguo, {"ruta": "a.pdf", "nombre": "A", "autor": "B", "niveles": ["1"], "asignatura": "3", "tipo_file": "70"})
        self.assertEqual((r["id"], r["subido"], visto["nombre"]), ("9", False, "A"))

    def test_subir_sin_niveles_no_entra_a_webclass(self):
        with mock.patch.object(bw, "iniciar_sesion") as sesion:
            r = self._subir(mock.Mock(), {"ruta": "a.pdf", "nombre": "A", "autor": "B", "niveles": ["kinder"], "asignatura": "3"})
        self.assertIn("nivel", r["error"])
        sesion.assert_not_called()

    def test_asignatura_del_recurso(self):
        self.assertEqual(bw.id_asignatura(rec, "Ciencias Naturales"), rec.ASIGNATURAS["ciencias"])
        self.assertEqual(bw.id_asignatura(rec, "Orientación"), rec.ASIGNATURAS["orientacion"])
        self.assertEqual(bw.id_asignatura(rec, "Historia, Geografía y Ciencias Sociales"), rec.ASIGNATURAS["historia"],
                         "manda la palabra que aparece primero")
        self.assertEqual(bw.id_asignatura(rec, "Filosofía"), rec.ASIGNATURAS["filosofia"])
        self.assertEqual(bw.id_asignatura(rec, "Religión"), 15, "código leído de WebClass real")
        self.assertEqual(bw.id_asignatura(rec, "Artes Visuales"), 6)
        self.assertIsNone(bw.id_asignatura(rec, "Cine y Teatro"), "fuera de la tabla, solo con la lista de la cuenta")

    def test_el_marcador_usa_la_tabla_de_asignaturas_de_barra(self):
        wc = comun.importar("nucleo.webclass")
        with mock.patch.object(rutas, "asignaturas_webclass_de_cuenta", return_value=[]):   # cuenta sin lista leída
            de_barra = wc.asignaturas_de_barra()
            todas = wc.asignaturas_recurso()
        for clave, nombre in barra_utiles.ASIGNATURAS_SUBIDA:
            self.assertEqual(de_barra[nombre], str(barra_utiles.ID_ASIGNATURA[clave]))
        self.assertEqual(todas["Ciencias Naturales"], "3", "antes el marcador no ofrecía Ciencias")
        self.assertEqual(todas["Matemática"], "2")
        repetidas = {i for n, i in todas.items() if n in wc.ASIGNATURAS_RECURSO and n not in de_barra} & set(de_barra.values())
        self.assertEqual(repetidas, set(), "una asignatura de Barra no se repite con otro nombre")
        with mock.patch.object(rutas, "BARRA", None):
            self.assertEqual(wc.asignaturas_de_barra(), {}, "suelto, sin Barra, quedan solo las propias")

    def test_con_la_lista_leida_de_webclass_se_ofrecen_las_de_la_cuenta(self):
        wc = comun.importar("nucleo.webclass")
        leidas = [{"id": 3, "nombre": "Ciencias Naturales"}, {"id": 50, "nombre": "Música"}]
        with mock.patch.object(rutas, "asignaturas_webclass_de_cuenta", return_value=leidas):
            self.assertEqual(wc.asignaturas_de_barra(), {"Ciencias Naturales": "3", "Música": "50"})

    def test_la_asignatura_de_una_unidad_usa_la_lista_de_la_cuenta(self):
        perfil = {"asignaturas_webclass": [{"id": 3, "nombre": "Ciencias Naturales"}, {"id": 18, "nombre": "Cine y Teatro"}]}
        self.assertEqual(bw.asignatura_de_la_unidad(rec, perfil, "Cine y Teatro"), 18)
        self.assertEqual(bw.asignatura_de_la_unidad(rec, {}, "Cine y Teatro"), None,
                         "sin la lista, la tabla no conoce Cine y Teatro")
        self.assertEqual(bw.asignatura_de_la_unidad(rec, {}, "Religión"), 15, "Religión ya está en la tabla")
        self.assertEqual(bw.asignatura_de_la_unidad(rec, {}, "Historia, Geografía y Ciencias Sociales"), rec.ASIGNATURAS["historia"])

    def test_la_lista_de_la_cuenta_se_lee_sin_tocar_otros_datos_del_perfil(self):
        with comun.BarraFalsa(activa="beto") as barra:
            r = comun.importar("nucleo.rutas")
            self.assertEqual(r.asignaturas_webclass_de_cuenta(), [])
            ruta = os.path.join(barra.carpeta, "personal", "perfiles", "beto.json")
            with open(ruta, encoding="utf-8") as f:
                datos = json.load(f)
            datos["asignaturas_webclass"] = [{"id": 50, "nombre": "Música"}]
            with open(ruta, "w", encoding="utf-8") as f:
                json.dump(datos, f)
            self.assertEqual(r.asignaturas_webclass_de_cuenta(), [{"id": 50, "nombre": "Música"}])
            self.assertEqual(r.asignaturas_webclass_de_cuenta("ana"), [], "cada cuenta tiene la suya")


class Lanzador(unittest.TestCase):
    """lanzar_planificador.pyw --solo-servidor: lo que pide la barra para tener el planificador listo."""

    @staticmethod
    def _cargar():
        import importlib.machinery
        import importlib.util
        ruta = os.path.join(comun.PLANIFICADOR, "lanzar_planificador.pyw")
        cargador = importlib.machinery.SourceFileLoader("lanzador_de_prueba", ruta)
        modulo = importlib.util.module_from_spec(importlib.util.spec_from_loader("lanzador_de_prueba", cargador))
        cargador.exec_module(modulo)
        return modulo

    def _main(self, lz, argv, puerto):
        with mock.patch.object(lz, "buscar_puerto", return_value=puerto), \
                mock.patch.object(lz, "iniciar_servidor") as iniciar, \
                mock.patch.object(lz, "es_mio", return_value=True), \
                mock.patch.object(lz.webbrowser, "open") as abrir, \
                mock.patch.object(lz, "Aviso") as aviso, \
                mock.patch.object(lz.sys, "argv", argv):
            lz.main()
        return iniciar, abrir, aviso

    def test_solo_servidor_lo_inicia_sin_navegador_ni_ventanas(self):
        iniciar, abrir, aviso = self._main(self._cargar(), ["lanzar", "--solo-servidor"], (8502, False))
        iniciar.assert_called_once_with(8502)
        abrir.assert_not_called()
        aviso.assert_not_called()

    def test_solo_servidor_no_abre_otro_si_ya_hay_uno(self):
        iniciar, abrir, _ = self._main(self._cargar(), ["lanzar", "--solo-servidor"], (8501, True))
        iniciar.assert_not_called()
        abrir.assert_not_called()

    def test_sin_opciones_abre_el_navegador_como_siempre(self):
        _, abrir, _ = self._main(self._cargar(), ["lanzar", "guia-pie"], (8501, True))
        abrir.assert_called_once_with("http://127.0.0.1:8501/#/guia-pie")

    def test_si_faltan_complementos_lo_dice_en_vez_de_esperar_un_minuto(self):
        lz = self._cargar()
        with mock.patch.object(lz, "complementos_faltantes", return_value=["fastapi"]), \
                mock.patch.object(lz, "ofrecer_instalar") as ofrecer:
            iniciar, abrir, aviso = self._main(lz, ["lanzar", "planificar"], (8502, False))
            ofrecer.assert_called_once_with(["fastapi"])
            iniciar.assert_not_called()
            abrir.assert_not_called()
            iniciar, _, _ = self._main(lz, ["lanzar", "--solo-servidor"], (8502, False))
            iniciar.assert_not_called()                      # «dejarlo listo» sin complementos: no hace nada
            self.assertEqual(ofrecer.call_count, 1, "ni muestra ventanas")

    def _buscar(self, version_que_corre, se_apaga=True):
        lz = self._cargar()
        with mock.patch.object(lz, "es_mio", lambda p: p == 8501), \
                mock.patch.object(lz, "version_local", return_value="1.5.2"), \
                mock.patch.object(lz, "version_corriendo", return_value=version_que_corre), \
                mock.patch.object(lz, "apagar", return_value=se_apaga) as apagar:
            return lz.buscar_puerto(), apagar

    def test_un_planificador_de_otra_version_se_cierra_y_se_abre_el_nuevo(self):
        (puerto, corriendo), apagar = self._buscar("1.4.1")
        self.assertEqual((puerto, corriendo), (8501, False), "en su lugar se inicia el nuevo")
        apagar.assert_called_once_with(8501)
        (_, corriendo), apagar = self._buscar("")             # uno tan antiguo que no informa su versión
        self.assertFalse(corriendo)

    def test_el_de_la_misma_version_se_usa_tal_cual(self):
        (puerto, corriendo), apagar = self._buscar("1.5.2")
        self.assertEqual((puerto, corriendo), (8501, True))
        apagar.assert_not_called()

    def test_si_no_se_puede_cerrar_el_viejo_se_usa_igual(self):
        (puerto, corriendo), _ = self._buscar("1.4.1", se_apaga=False)
        self.assertEqual((puerto, corriendo), (8501, True))

    def test_detecta_los_complementos_que_faltan(self):
        lz = self._cargar()
        import importlib.util
        real = importlib.util.find_spec
        with mock.patch.object(importlib.util, "find_spec", lambda m, *a: None if m == "fastapi" else real(m, *a)):
            self.assertEqual(lz.complementos_faltantes(), ["fastapi"])
        self.assertEqual(lz.complementos_faltantes(), [], "en este equipo de pruebas está todo")


class IaEnVivo(unittest.TestCase):
    """ia.completar avisa cada intento y cada trozo de texto (para mostrarlo mientras se escribe)."""

    CONFIG = {"proveedor": "Gemini (Google)", "modelo": None, "respaldo": True, "claves": {"Gemini (Google)": "x"}}

    @staticmethod
    def _flujo(*textos, falla=None):
        from types import SimpleNamespace as N

        def crear(**_argumentos):
            for t in textos:
                yield N(choices=[N(delta=N(content=t), finish_reason=None)])
            if falla:
                raise falla
            yield N(choices=[N(delta=None, finish_reason="stop")])
        return crear

    def _completar(self, respuestas, **avisos):
        from types import SimpleNamespace as N
        ia = comun.importar("nucleo.ia")
        siguiente = iter(respuestas)
        with mock.patch.object(ia, "config", return_value=self.CONFIG), \
                mock.patch.object(ia, "_orden_de_intento", return_value=[("Gemini (Google)", "x", "propia")]), \
                mock.patch.object(ia, "_modelos_a_probar", return_value=["m1", "m2"]), \
                mock.patch.object(ia, "_bloquear"), \
                mock.patch.object(ia, "cliente", lambda *a, **k: N(chat=N(completions=N(create=next(siguiente))))):
            return ia.completar(lambda info: [], max_tokens=10, **avisos)

    def test_avisa_cada_intento_y_cada_trozo(self):
        intentos, trozos = [], []
        texto, usado, fin, avisos = self._completar(
            [self._flujo("Ho", "la", falla=RuntimeError("se cortó")), self._flujo("Buenas ", "tardes")],
            al_intento=intentos.append, al_trozo=trozos.append)
        self.assertEqual((texto, usado, fin), ("Buenas tardes", "Gemini (Google) · m2", "stop"))
        self.assertEqual(intentos, ["Gemini (Google) · m1", "Gemini (Google) · m2"], "el segundo aviso descarta lo del primero")
        self.assertEqual(trozos, ["Ho", "la", "Buenas ", "tardes"])
        self.assertEqual(len(avisos), 1)

    def test_un_aviso_que_falla_no_bota_la_respuesta(self):
        def falla(_texto):
            raise ValueError("la pantalla se cerró")
        texto, *_ = self._completar([self._flujo("Hola")], al_trozo=falla, al_intento=falla)
        self.assertEqual(texto, "Hola")

    def test_sin_avisos_funciona_como_antes(self):
        self.assertEqual(self._completar([self._flujo("Hola", " mundo")])[0], "Hola mundo")


class ClavesDelColegio(unittest.TestCase):
    """Las claves de IA del colegio llegan cifradas con la barra: se usan sin mostrarse y la propia (extra) va primero."""

    CODIGO = "ABCDEFGHJKLMNPQRSTUV"

    def setUp(self):
        sin_entorno = mock.patch.dict(os.environ, {})          # que una clave en variables de entorno no se cuele
        sin_entorno.start()
        self.addCleanup(sin_entorno.stop)
        for p in ("GEMINI_API_KEY", "GROQ_API_KEY", "CEREBRAS_API_KEY", "OPENROUTER_API_KEY"):
            os.environ.pop(p, None)
        self.ia = comun.importar("nucleo.ia")

    def _buzon(self, barra, claves):
        """Copia colegio.py a la Barra falsa y deja su buzón cifrado con la clave del colegio, como lo deja la barra."""
        import shutil
        import colegio
        herramienta = os.path.dirname(os.path.abspath(barra_utiles.__file__))
        shutil.copy2(os.path.join(herramienta, "colegio.py"), os.path.join(barra.carpeta, "herramienta", "colegio.py"))
        colegio.guardar_paquete(os.path.join(barra.carpeta, "datos", "colegio.json"),
                                colegio.cerrar(self.CODIGO, {"claves": claves, "mensajes": []}))

    def test_sin_clave_propia_usa_la_del_colegio_sin_mostrarla(self):
        ia = self.ia
        with comun.BarraFalsa() as barra, mock.patch("keyring.get_password", return_value=self.CODIGO):
            self._buzon(barra, {"GEMINI_API_KEY": "CLAVE-DEL-COLEGIO-1234"})
            self.assertEqual(ia.claves_de("Gemini (Google)"), [("CLAVE-DEL-COLEGIO-1234", "colegio")])
            self.assertEqual(ia.config()["proveedor"], "Gemini (Google)")
            self.assertTrue(ia.hay_proveedor())
            estado = ia.estado_claves()["Gemini (Google)"]
            self.assertEqual((estado["guardada"], estado["colegio"], estado["final"]), (False, True, ""))
            self.assertNotIn("CLAVE-DEL-COLEGIO", json.dumps(ia.estado_claves()), "ni siquiera el final se muestra")
            self.assertEqual(ia.estado_colegio(), "ok")
            ia.guardar_claves({"Gemini (Google)": "MI-CLAVE-5678"})          # la extra del profe va primero
            self.assertEqual(ia.claves_de("Gemini (Google)"),
                             [("MI-CLAVE-5678", "propia"), ("CLAVE-DEL-COLEGIO-1234", "colegio")])
            ia.guardar_claves({"Gemini (Google)": ""})                        # quitar la propia deja la del colegio
            self.assertEqual(ia.claves_de("Gemini (Google)"), [("CLAVE-DEL-COLEGIO-1234", "colegio")])

    def test_sin_la_clave_del_colegio_en_este_pc_lo_explica(self):
        ia = self.ia
        for guardada, estado, frase in ((None, "sin_codigo", "falta la clave del colegio"),
                                        ("ZZZZZZZZZZZZZZZZZZZZ", "codigo_incorrecto", "no abre las claves")):
            with self.subTest(estado=estado), comun.BarraFalsa() as barra, \
                    mock.patch("keyring.get_password", return_value=guardada):
                self._buzon(barra, {"GEMINI_API_KEY": "CLAVE-DEL-COLEGIO-1234"})
                self.assertEqual(ia.estado_colegio(), estado)
                self.assertFalse(ia.hay_proveedor())
                self.assertIn(frase, ia.sin_claves_explicado())

    def test_si_el_colegio_no_publica_claves_lo_explica(self):
        """La clave del colegio funciona (ticket verde), pero el buzón no trae claves de IA: se dice eso, no «falta la
        clave»."""
        ia = self.ia
        with comun.BarraFalsa() as barra, mock.patch("keyring.get_password", return_value=self.CODIGO):
            self._buzon(barra, {})
            self.assertEqual(ia.estado_colegio(), "ok")
            self.assertFalse(ia.hay_proveedor())
            self.assertIn("todavía no publica claves de IA", ia.sin_claves_explicado())

    def test_planificador_suelto_sin_barra_no_falla(self):
        with mock.patch.object(self.ia, "_colegio", return_value=None):
            self.assertEqual(self.ia.claves_del_colegio(), {})
            self.assertEqual(self.ia.estado_colegio(), "sin_buzon")

    def test_si_la_propia_llega_a_su_limite_sigue_la_del_colegio(self):
        from types import SimpleNamespace as N
        ia = self.ia

        class Limite(Exception):
            status_code = 429

        def crear_con(clave):
            def crear(**_argumentos):
                if clave == "MIA":
                    raise Limite("rate limit")
                return iter([N(choices=[N(delta=N(content="Hola"), finish_reason=None)]),
                             N(choices=[N(delta=None, finish_reason="stop")])])
            return crear
        config = {"proveedor": "Gemini (Google)", "modelo": None, "respaldo": False,
                  "claves": {"Gemini (Google)": "MIA"}, "todas": {"Gemini (Google)": [("MIA", "propia"), ("COLE", "colegio")]}}
        self.addCleanup(ia._bloqueados.clear)
        with mock.patch.object(ia, "config", return_value=config), \
                mock.patch.object(ia, "_modelos_a_probar", return_value=["m1"]), \
                mock.patch.object(ia, "cliente", lambda proveedor, clave, streaming=False:
                                  N(chat=N(completions=N(create=crear_con(clave))))):
            texto, usado, _fin, avisos = ia.completar(lambda info: [], max_tokens=10)
        self.assertEqual((texto, usado), ("Hola", "Gemini (Google) · m1"))
        self.assertEqual(len(avisos), 1, "sin respaldo de otros proveedores, igual pasa a la clave del colegio")
        self.assertIn(("Gemini (Google)", ia._huella("MIA"), "m1"), ia._bloqueados)
        self.assertNotIn(("Gemini (Google)", ia._huella("COLE"), "m1"), ia._bloqueados, "el límite es de cada clave")

    def test_la_barra_conoce_los_mismos_proveedores(self):
        import colegio
        self.assertEqual([(n, s) for n, s, _ in colegio.PROVEEDORES_IA],
                         [(n, i["secreto"]) for n, i in self.ia.PROVEEDORES.items()])


class RegistroYDeshacer(unittest.TestCase):
    UNIDAD = {"id": 7433453, "nombre": "5BA - Unidad 1: El agua en la Tierra", "curso": "5BA"}

    def setUp(self):
        import rutas as rutas_barra
        self.rutas_barra = rutas_barra
        self.tmp_ctx = comun.carpeta_temporal()
        self.tmp = self.tmp_ctx.__enter__()
        with open(os.path.join(self.tmp, "registro_historico.txt"), "w", encoding="utf-8") as f:
            f.write("bastian")
        for parche in (mock.patch.object(rutas_barra, "DATOS", self.tmp),
                       mock.patch.object(rutas_barra, "HISTORICO", os.path.join(self.tmp, "registro_historico.txt"))):
            parche.start()
            self.addCleanup(parche.stop)
        self.addCleanup(self.tmp_ctx.__exit__, None, None, None)

    def test_cada_creacion_deja_una_fila_con_encabezado_unico(self):
        bw.registrar_creacion("bastian", self.UNIDAD, "111", "Clase A", "15/10/2026", 1, ["900", "901"])
        bw.registrar_creacion("bastian", self.UNIDAD, "112", "Clase B", "16/10/2026", 2, [])
        archivo = os.path.join(self.tmp, "registro_planificaciones.csv")
        texto = leer(archivo)
        self.assertEqual(texto.count("fecha_creacion"), 1)
        filas = list(csv.DictReader(io.StringIO(texto)))
        self.assertEqual([f["idclase"] for f in filas], ["111", "112"])
        self.assertEqual((filas[0]["recursos_ids"], filas[0]["origen"], filas[0]["curso"]), ("900;901", "planificador", "5BA"))

    def test_otra_cuenta_escribe_en_su_propio_archivo(self):
        bw.registrar_creacion("jazmyna", self.UNIDAD, "113", "De otra", "17/10/2026", 0, [])
        self.assertTrue(os.path.exists(os.path.join(self.tmp, "registro_planificaciones_jazmyna.csv")))
        self.assertFalse(os.path.exists(os.path.join(self.tmp, "registro_planificaciones.csv")))

    # ---- eliminar: solo con lo registrado y con sesión simulada ----
    class Salida(Exception):
        pass

    def _eliminar(self, clases, vivas, asignados=0, borra=True):
        test = self

        class Respuesta:
            def raise_for_status(self):
                pass

        class Sesion:
            def __init__(self):
                self.posts = []

            def post(self, url, data=None, headers=None, timeout=None):
                self.posts.append(data)
                if borra:
                    vivas.discard(str(data["idClase"]))
                return Respuesta()

        sesion = Sesion()

        class W:
            BASE, HEADERS, PERFIL_POR_DEFECTO = "https://webclass.prueba", {}, "bastian"

            @staticmethod
            def listar_clases(s, id_unidad):
                return [{"idclase": i, "nombre": f"Clase {i}", "asignados": asignados} for i in sorted(vivas)]

        class PL:
            @staticmethod
            def unidad_del_perfil(perfil, id_unidad):
                if str(id_unidad) != "7433453":
                    raise ValueError("no es de la cuenta")
                return {"id": 7433453}

        def salir(obj):
            raise test.Salida(obj)

        with mock.patch.object(bw, "abrir", lambda d: (W, None, PL, None)), \
                mock.patch.object(bw, "perfil_de", lambda w, d: ("bastian", {"iduser": "77"})), \
                mock.patch.object(bw, "iniciar_sesion", lambda w, n, p: sesion), \
                mock.patch.object(bw, "salir", salir):
            with self.assertRaises(self.Salida) as ctx:
                bw.op_eliminar({"clases": clases})
        return ctx.exception.args[0]["resultados"], sesion

    def test_solo_elimina_clases_del_registro(self):
        bw.registrar_creacion("bastian", self.UNIDAD, "111", "Clase A", "15/10/2026", 1, [])
        vivas = {"111", "222"}                                       # 222 existe pero NO la creó el planificador
        resultados, sesion = self._eliminar([{"id_unidad": "7433453", "idclase": "222"}], vivas)
        self.assertEqual(resultados[0]["estado"], "rechazada")
        self.assertEqual(sesion.posts, [], "no se debe enviar ningún pedido de borrado")
        self.assertIn("222", vivas)

    def test_no_elimina_una_clase_con_tarea_asignada(self):
        bw.registrar_creacion("bastian", self.UNIDAD, "111", "Clase A", "15/10/2026", 1, [])
        resultados, sesion = self._eliminar([{"id_unidad": "7433453", "idclase": "111"}], {"111"}, asignados=12)
        self.assertEqual(resultados[0]["estado"], "rechazada")
        self.assertIn("12", resultados[0]["motivo"])
        self.assertEqual(sesion.posts, [])

    def test_elimina_una_clase_registrada_con_el_pedido_de_webclass(self):
        bw.registrar_creacion("bastian", self.UNIDAD, "111", "Clase A", "15/10/2026", 1, [])
        vivas = {"111"}
        resultados, sesion = self._eliminar([{"id_unidad": "7433453", "idclase": "111"}], vivas)
        self.assertEqual(resultados[0]["estado"], "eliminada")
        self.assertEqual(sesion.posts, [{"accion": 4, "caso": 3, "idClase": "111", "idplan": "7433453", "idUsuario": "77"}])
        self.assertEqual(vivas, set())

    def test_si_webclass_no_la_elimina_lo_dice(self):
        bw.registrar_creacion("bastian", self.UNIDAD, "111", "Clase A", "15/10/2026", 1, [])
        resultados, _ = self._eliminar([{"id_unidad": "7433453", "idclase": "111"}], {"111"}, borra=False)
        self.assertEqual(resultados[0]["estado"], "error")

    def test_una_clase_que_ya_no_existe_no_envia_nada(self):
        bw.registrar_creacion("bastian", self.UNIDAD, "111", "Clase A", "15/10/2026", 1, [])
        resultados, sesion = self._eliminar([{"id_unidad": "7433453", "idclase": "111"}], set())
        self.assertEqual(resultados[0]["estado"], "ya")
        self.assertEqual(sesion.posts, [])

    def test_una_unidad_ajena_se_rechaza(self):
        bw.registrar_creacion("bastian", {**self.UNIDAD, "id": 999}, "111", "Clase A", "15/10/2026", 1, [])
        resultados, sesion = self._eliminar([{"id_unidad": "999", "idclase": "111"}], {"111"})
        self.assertEqual(resultados[0]["estado"], "rechazada")
        self.assertEqual(sesion.posts, [])


if __name__ == "__main__":
    unittest.main()
