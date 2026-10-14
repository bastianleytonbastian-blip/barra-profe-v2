"""Planificación en Excel: datos de la clase y una fila por sección."""
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font, PatternFill

from nucleo.webclass import partes_con_formato


def texto_enriquecido(texto):
    """Texto de un cuadro -> celda de Excel con subtítulos y rótulos en negrita (mismo criterio que WebClass)."""
    bloques = []
    negrita = InlineFont(b=True)
    for parte in partes_con_formato(texto):
        tipo = parte[0]
        if tipo == "titulo":
            bloques.append(TextBlock(negrita, parte[1] + "\n"))
        elif tipo == "rotulo":
            bloques.append(TextBlock(negrita, parte[1] + " "))
            bloques.append(parte[2] + "\n")
        elif tipo == "vineta":
            bloques.append("• " + parte[1] + "\n")
        elif tipo == "parrafo":
            bloques.append(parte[1] + "\n")
        else:
            bloques.append("\n")
    if bloques and isinstance(bloques[-1], str):
        bloques[-1] = bloques[-1].rstrip("\n") or " "
    return CellRichText(bloques) if bloques else ""


def generar_excel(meta, filas_plan):
    """meta: {asignatura, curso, nombre, tiempo, tema, fecha}; filas_plan: [(sección, texto)]. Devuelve bytes .xlsx."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Planificación"
    negrita = Font(bold=True)
    ajuste = Alignment(wrap_text=True, vertical="top")

    filas_info = [
        ("Asignatura", meta.get("asignatura", "")),
        ("Curso", meta.get("curso", "")),
        ("Nombre de la clase", meta.get("nombre", "")),
        ("Duración", meta.get("tiempo", "")),
        ("Contenido / tema", meta.get("tema", "")),
        ("Fecha", meta.get("fecha", date.today().strftime("%d/%m/%Y"))),
    ]
    for i, (etiqueta, valor) in enumerate(filas_info, start=1):
        hoja.cell(row=i, column=1, value=etiqueta).font = negrita
        hoja.cell(row=i, column=2, value=valor).alignment = ajuste

    fila = len(filas_info) + 2
    relleno = PatternFill("solid", fgColor="D9E1F2")
    for col, titulo in ((1, "Sección"), (2, "Planificación")):
        celda = hoja.cell(row=fila, column=col, value=titulo)
        celda.font = negrita
        celda.fill = relleno

    for nombre, texto in filas_plan:
        fila += 1
        celda_nombre = hoja.cell(row=fila, column=1, value=nombre)
        celda_nombre.font = negrita
        celda_nombre.alignment = ajuste
        hoja.cell(row=fila, column=2, value=texto_enriquecido(texto)).alignment = ajuste
        # Alto aproximado según la cantidad de texto (Excel no lo ajusta solo al abrir)
        lineas = sum(max(1, len(parrafo) // 110 + 1) for parrafo in (texto or "").split("\n"))
        hoja.row_dimensions[fila].height = min(409, 15 * lineas)

    hoja.column_dimensions["A"].width = 22
    hoja.column_dimensions["B"].width = 120
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
