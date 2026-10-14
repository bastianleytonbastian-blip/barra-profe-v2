# -*- coding: utf-8 -*-
"""Guía PIE en Word (.docx) con el mismo contenido y reglas que la versión HTML (armar_guia_pie.estructura_guia).
Los pictogramas (ARASAAC) y las fotos de Wikimedia se descargan al crear el archivo."""
import json
import os
import re
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from nucleo import armar_guia_pie as ag
from nucleo.armar_guia_pie import g

ANCHO = 18.6  # cm útiles en A4 con márgenes de 1,2 cm
COLOR = {"rojo": "CC0000", "verde": "2E7D32", "gris": "AAAAAA", "azul": "1787C8", "navy": "1F3B57",
         "barra": "DCE6F1", "espacio": "CCCCCC", "fondo": "F5F5F5"}
_pictos, _fotos = {}, {}


# ---------- Descargas ----------

def _bajar(url, json_=False):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (Planificador escolar)"})
    with urllib.request.urlopen(req, timeout=20) as r:
        datos = r.read()
    return json.loads(datos) if json_ else datos


def _variantes(q):
    """Igual que en la guía HTML: plural -> singular, verbo conjugado -> infinitivo."""
    s = "".join(c for c in __import__("unicodedata").normalize("NFD", q) if not __import__("unicodedata").combining(c))
    v = [q] + ([s] if s != q else [])
    for patron, rep in [(r"es$", ""), (r"s$", ""), (r"aron$", "ar"), (r"ieron$", "er"), (r"ieron$", "ir"), (r"aban?$", "ar"),
                        (r"ian?$", "er"), (r"ian?$", "ir"), (r"io$", "er"), (r"io$", "ir"), (r"o$", "ar"), (r"an$", "ar"),
                        (r"en$", "er"), (r"en$", "ir")]:
        if re.search(patron, s):
            c = re.sub(patron, rep, s)
            if len(c) >= 3 and c not in v:
                v.append(c)
    return v[:5]


def picto(termino):
    """Bytes PNG del pictograma ARASAAC para un término (o None)."""
    if termino in _pictos:
        return _pictos[termino]
    resultado = None
    for q in _variantes(termino):
        try:
            j = _bajar("https://api.arasaac.org/v1/pictograms/es/search/" + urllib.parse.quote(q), json_=True)
            if j:
                pid = j[0]["_id"]
                resultado = _bajar(f"https://static.arasaac.org/pictograms/{pid}/{pid}_300.png")
                break
        except Exception:
            continue
    _pictos[termino] = resultado
    return resultado


def foto_web(busqueda):
    """Primera foto de Wikimedia Commons para la búsqueda (bytes) o None."""
    if busqueda in _fotos:
        return _fotos[busqueda]
    resultado = None
    try:
        url = ("https://commons.wikimedia.org/w/api.php?action=query&generator=search&gsrnamespace=6&gsrlimit=10"
               "&gsrsearch=" + urllib.parse.quote("filetype:bitmap " + busqueda) +
               "&prop=imageinfo&iiprop=url|mime&iiurlwidth=800&format=json")
        paginas = (_bajar(url, json_=True).get("query") or {}).get("pages") or {}
        for p in sorted(paginas.values(), key=lambda p: p.get("index", 0)):
            info = (p.get("imageinfo") or [{}])[0]
            if re.search(r"jpeg|png", info.get("mime", "")):
                resultado = _bajar(info.get("thumburl") or info["url"])
                break
    except Exception:
        pass
    _fotos[busqueda] = resultado
    return resultado


def _precargar(secciones, objetivo):
    """Descarga en paralelo todos los pictogramas y fotos que usará la guía."""
    terminos, busquedas = set(), set()
    textos = [objetivo]
    for _, _, bloques in secciones:
        for b in bloques:
            if b[0] in ("escribir", "copiar", "completar"):
                textos += b[2]
            elif b[0] == "marcar":
                textos += [b[2]] + b[3]
            elif b[0] in ("reconocer", "unir"):
                textos += [c for _, c in b[2]]
            elif b[0] == "web":
                busquedas.add(b[1])
    for t in textos:
        for w in t.split():
            termino = g.termino_picto(w)
            if termino:
                terminos.add(termino)
    with ThreadPoolExecutor(8) as ex:
        list(ex.map(picto, terminos))
        list(ex.map(foto_web, busquedas))


# ---------- Formato de tablas y celdas ----------

# Orden obligatorio de los elementos dentro de <w:tcPr> (Word marca el archivo como dañado si no se respeta)
_ORDEN_TCPR = ["tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection",
               "tcFitText", "vAlign", "hideMark"]


def _poner_en_tcpr(celda, elemento):
    tcPr = celda._tc.get_or_add_tcPr()
    nombre = elemento.tag.split("}")[1]
    for viejo in tcPr.findall(qn(f"w:{nombre}")):
        tcPr.remove(viejo)
    posicion = _ORDEN_TCPR.index(nombre)
    for i, hijo in enumerate(list(tcPr)):
        etiqueta = hijo.tag.split("}")[1]
        if etiqueta in _ORDEN_TCPR and _ORDEN_TCPR.index(etiqueta) > posicion:
            tcPr.insert(i, elemento)
            return
    tcPr.append(elemento)


def _sombra(celda, color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    _poner_en_tcpr(celda, shd)


def _bordes(celda, visibles=True, grosor=8):
    bordes = OxmlElement("w:tcBorders")
    for lado in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single" if visibles else "nil")
        if visibles:
            e.set(qn("w:sz"), str(grosor)); e.set(qn("w:color"), "000000")
        bordes.append(e)
    _poner_en_tcpr(celda, bordes)


def _alto(fila, cm):
    trPr = fila._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(cm * 567))); h.set(qn("w:hRule"), "atLeast")
    trPr.append(h)


def _tabla(doc, filas, columnas, anchos=None):
    t = doc.add_table(rows=filas, cols=columnas)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    anchos = anchos or [ANCHO / columnas] * columnas
    for fila in t.rows:
        for celda, ancho in zip(fila.cells, anchos):
            celda.width = Cm(ancho)
    return t


def _texto(celda_o_parrafo, texto, tam=24, color=None, centrado=False, negrita=True):
    p = celda_o_parrafo.paragraphs[0] if hasattr(celda_o_parrafo, "paragraphs") else celda_o_parrafo
    if centrado:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(texto)
    r.font.size, r.font.bold, r.font.name = Pt(tam), negrita, "Calibri"
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return p


def _imagen(parrafo, datos, alto_cm=None, ancho_cm=None):
    if not datos:
        return
    try:
        parrafo.add_run().add_picture(BytesIO(datos), height=Cm(alto_cm) if alto_cm else None,
                                      width=Cm(ancho_cm) if ancho_cm else None)
    except Exception:
        pass


def _espacio(doc, pt=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(" "); r.font.size = Pt(pt)


# ---------- Bloques de la guía ----------

def _cuadricula(doc, palabras_linea, modo, guias=None, verbo_rojo=False):
    """Una línea de cuadrícula: fila de pictogramas, fila modelo y fila para escribir (copy / guide)."""
    n = len(palabras_linea)
    t = _tabla(doc, 3, n)
    # Fila de pictogramas: una celda combinada por palabra
    i = 0
    while i < n:
        if palabras_linea[i]["type"] == "letter":
            j = i
            while j < n and palabras_linea[j]["type"] == "letter":
                j += 1
            palabra = "".join(x["ch"] for x in palabras_linea[i:j])
            celda = t.cell(0, i).merge(t.cell(0, j - 1)) if j - 1 > i else t.cell(0, i)
            nw = g.norm(palabra)
            datos = None
            if nw in g.ARTS and os.path.exists(g.art_picto(nw)):
                datos = open(g.art_picto(nw), "rb").read()
            elif g.termino_picto(palabra):
                datos = picto(g.termino_picto(palabra))
            celda.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _imagen(celda.paragraphs[0], datos, alto_cm=1.0)
            i = j
        else:
            i += 1
    for c in t.rows[0].cells:
        _bordes(c, False)
    for k, celda in enumerate(palabras_linea):
        arriba, abajo = t.cell(1, k), t.cell(2, k)
        for c in (arriba, abajo):
            _bordes(c, True)
        if celda["type"] == "space":
            _sombra(arriba, COLOR["espacio"]); _sombra(abajo, COLOR["espacio"])
            continue
        color = COLOR["rojo"] if celda.get("_verbo") else (COLOR["verde"] if celda["is_art"] else "000000")
        _texto(arriba, celda["ch"], 24, color, centrado=True)
        if modo == "guide" and guias and k in guias:
            _texto(abajo, celda["ch"], 24, COLOR["gris"], centrado=True)
    _alto(t.rows[1], 1.3); _alto(t.rows[2], 1.3)
    _espacio(doc, 3)


def _frases_cuadricula(doc, frases, modo):
    for frase in frases:
        for linea in g._wrap_lines(frase):
            guias = None
            if modo == "guide":
                for k, c in enumerate(linea):
                    c["_idx"] = k
                guias = g._guides_for_line(linea)
            _cuadricula(doc, linea, modo, guias)
        _espacio(doc, 6)


def _objetivo(doc, objetivo):
    _cuadricula(doc, [{"type": "letter", "ch": c, "is_art": False} for c in "OBJETIVO:"], "copy")
    palabras = objetivo.split(" ")
    celdas = []
    for wi, w in enumerate(palabras):
        art = g.norm(w) in g.ARTS
        celdas += [{"type": "letter", "ch": ch, "is_art": art, "_verbo": wi == 0} for ch in w]
        celdas.append({"type": "space"})
    celdas.pop()
    lineas, actual = [], []
    for c in celdas:
        actual.append(c)
        if len(actual) >= g.MAX_CELLS and c["type"] == "space":
            lineas.append(actual[:-1]); actual = []
    if actual:
        lineas.append(actual)
    for linea in lineas:
        _cuadricula(doc, linea, "copy")


def _instruccion(doc, texto):
    """Instrucción con el verbo en rojo y su pictograma, artículos en verde con el suyo."""
    palabras = texto.split(" ")
    t = _tabla(doc, 2, len(palabras), [min(ANCHO / len(palabras), 4.5)] * len(palabras))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, w in enumerate(palabras):
        base, ruta, color = g.norm(w), None, "000000"
        if i == 0:
            ruta, color = g.verbo_picto(base), COLOR["rojo"]
        elif base in g.ARTS:
            ruta, color = g.art_picto(base), COLOR["verde"]
        for c in (t.cell(0, i), t.cell(1, i)):
            _bordes(c, False)
        if ruta and os.path.exists(ruta):
            t.cell(0, i).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _imagen(t.cell(0, i).paragraphs[0], open(ruta, "rb").read(), alto_cm=1.1)
        _texto(t.cell(1, i), w, 20, color, centrado=True)
    _espacio(doc, 4)


def _picto_frase(parrafo, texto, alto=1.1):
    palabras = texto.split()
    termino = next((g.PICTO_MAP[g.norm(w)] for w in palabras if g.norm(w) in g.PICTO_MAP), None) or \
        next((x for x in (g.termino_picto(w) for w in palabras) if x), None)
    if termino:
        _imagen(parrafo, picto(termino), alto_cm=alto)


def _ilustracion(doc, datos, pie):
    if datos:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _imagen(p, datos, ancho_cm=10)
    if pie:
        _texto(doc.add_paragraph(), pie, 12, "333333", centrado=True, negrita=True)


def _marcar(doc, instr, pregunta, opciones):
    _instruccion(doc, instr)
    p = doc.add_paragraph(); _picto_frase(p, pregunta); _texto(p, " " + pregunta, 18)
    t = _tabla(doc, len(opciones), 4, [1.4, 1.2, 2.0, ANCHO - 4.6])
    for i, op in enumerate(opciones):
        celdas = t.rows[i].cells
        for c in celdas:
            _bordes(c, False)
        _texto(celdas[0], "○", 28, centrado=True, negrita=False)
        _texto(celdas[1], "ABCDE"[i], 20, COLOR["azul"], centrado=True)
        celdas[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _picto_frase(celdas[2].paragraphs[0], op)
        _texto(celdas[3], op, 18)
        _alto(t.rows[i], 1.4)
    _espacio(doc, 8)


def _reconocer(doc, instr, pares):
    _instruccion(doc, instr)
    t = _tabla(doc, 1, len(pares))
    for celda, (datos, concepto) in zip(t.rows[0].cells, pares):
        _bordes(celda, True, 16)
        p = celda.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _imagen(p, datos, ancho_cm=min(5.0, ANCHO / len(pares) - 0.6))
        p2 = celda.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _picto_frase(p2, concepto, 1.0)
        _texto(celda.add_paragraph(), concepto, 14, centrado=True)
    _espacio(doc, 8)


def _unir(doc, instr, pares):
    _instruccion(doc, instr)
    n = len(pares)
    conceptos = [pares[(i + 1) % n][1] for i in range(n)]  # mismo orden barajado que la guía HTML
    t = _tabla(doc, n, 5, [1.2, 5.5, 1.5, 1.2, ANCHO - 9.4])
    for i in range(n):
        num, img, hueco, letra, conc = t.rows[i].cells
        for c, visible in ((num, True), (img, True), (hueco, False), (letra, True), (conc, True)):
            _bordes(c, visible, 16)
        _sombra(num, COLOR["azul"]); _texto(num, str(i + 1), 20, "FFFFFF", centrado=True)
        img.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _imagen(img.paragraphs[0], pares[i][0], ancho_cm=5.0)
        _sombra(letra, COLOR["rojo"]); _texto(letra, "ABCDE"[i], 20, "FFFFFF", centrado=True)
        _picto_frase(conc.paragraphs[0], conceptos[i]); _texto(conc, " " + conceptos[i], 18)
    _espacio(doc, 8)


def _recuadro(doc, texto=None, alto_cm=None):
    t = _tabla(doc, 1, 1)
    c = t.cell(0, 0)
    _bordes(c, True, 24)
    if texto:
        _sombra(c, COLOR["fondo"])
        for i, linea in enumerate(texto.split("\n")):
            _texto(c if i == 0 else c.add_paragraph(), linea, 18)
    if alto_cm:
        _alto(t.rows[0], alto_cm)


def _encabezado_seccion(doc, numero, titulo, primera):
    if not primera:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    t = _tabla(doc, 1, 2, [1.6, ANCHO - 1.6])
    num, tit = t.rows[0].cells
    for c in (num, tit):
        _bordes(c, False)
    _sombra(num, COLOR["navy"]); _texto(num, str(numero), 20, "FFFFFF", centrado=True)
    _sombra(tit, COLOR["barra"]); _texto(tit, titulo, 18, COLOR["navy"])
    _espacio(doc, 6)


def _encabezado(doc, nivel, enc=None):
    colegio, depto, linea3 = g.lineas_encabezado(nivel, enc)
    ruta = (enc or {}).get("logo")
    tiene_logo = bool(ruta and os.path.exists(ruta))
    t = _tabla(doc, 1, 2, [2.2, ANCHO - 2.2])
    logo, textos = t.rows[0].cells
    for c in (logo, textos):
        _bordes(c, False)
    if tiene_logo:
        with open(ruta, "rb") as f:
            _imagen(logo.paragraphs[0], f.read(), ancho_cm=1.9)
    lineas = [(colegio, 14), (depto, 10.5), (linea3, 10.5)]
    primera = True
    for texto, tam in lineas:
        if not texto:
            continue
        _texto(textos if primera else textos.add_paragraph(), texto, tam)
        primera = False
    f = _tabla(doc, 2, 4, [4.0, 2.0, 2.0, 2.0])
    f.alignment = WD_TABLE_ALIGNMENT.CENTER
    for fila in f.rows:
        for c in fila.cells:
            _bordes(c, True)
        _alto(fila, 1.0)
    _texto(f.cell(0, 0), "FECHA", 20, centrado=True)
    _espacio(doc, 8)


def guia_a_word(spec, nivel, imagenes, encabezado=None):
    """spec: contenido de la guía; nivel: 'Sextos', 'Séptimos', 'Primeros Medios'…;
    imagenes: {id: (bytes, ext)}. Devuelve (bytes_docx, avisos)."""
    avisos = []
    objetivo, secciones = ag.estructura_guia(spec, list(imagenes), avisos, (encabezado or {}).get("asignatura", ""))
    _precargar(secciones, objetivo)

    doc = Document()
    for s in doc.sections:
        s.page_height, s.page_width = Cm(29.7), Cm(21)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(1.2)
    estilo = doc.styles["Normal"]
    estilo.font.name, estilo.font.size, estilo.font.bold = "Calibri", Pt(24), True
    estilo.paragraph_format.space_after = Pt(2)

    _encabezado(doc, nivel, encabezado)
    _objetivo(doc, objetivo)
    datos = lambda i: imagenes[i][0] if i in imagenes else None
    for k, (numero, titulo, bloques) in enumerate(secciones):
        _encabezado_seccion(doc, numero, titulo, primera=False)
        for b in bloques:
            tipo = b[0]
            if tipo == "ilus":
                _ilustracion(doc, datos(b[1]), b[2])
            elif tipo == "web":
                foto = foto_web(b[1])
                _ilustracion(doc, foto, b[2] if foto else f"[IMAGEN: {b[2]}]")
            elif tipo in ("escribir", "copiar"):
                _instruccion(doc, b[1]); _frases_cuadricula(doc, b[2], "copy")
            elif tipo == "completar":
                _instruccion(doc, b[1]); _frases_cuadricula(doc, b[2], "guide")
            elif tipo == "marcar":
                _marcar(doc, b[1], b[2], b[3])
            elif tipo == "conversar":
                _recuadro(doc, "PREGUNTA PARA CONVERSAR EN CLASE:\n" + b[1])
            elif tipo == "reconocer":
                _reconocer(doc, b[1], [(datos(i), c) for i, c in b[2]])
            elif tipo == "unir":
                _unir(doc, b[1], [(datos(i), c) for i, c in b[2]])
            elif tipo == "dibujo":
                _instruccion(doc, ag.instruccion_valida(b[1], "DIBUJAR")); _recuadro(doc, alto_cm=15)
    salida = BytesIO()
    doc.save(salida)
    return salida.getvalue(), avisos
