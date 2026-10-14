"""Guías PIE: lectura de material, pedido a la IA, reparación del JSON y revisión ortográfica. Sin Streamlit."""
import difflib
import hashlib
import json
import re
import unicodedata
from io import BytesIO

from docx import Document
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader

from nucleo import ia

_ORDINALES = ["Primeros", "Segundos", "Terceros", "Cuartos", "Quintos", "Sextos", "Séptimos", "Octavos"]
NIVELES = {f"{n}° Básico": _ORDINALES[n - 1] for n in range(1, 9)}
NIVELES.update({f"{n}° Medio": f"{_ORDINALES[n - 1]} Medios" for n in range(1, 5)})
LADO_MINIMO = 120      # px: descarta íconos, viñetas y líneas decorativas
LADO_MAXIMO = 1000     # px: reduce fotos grandes para que el HTML no pese demasiado
MAX_IMAGENES = 60
MAX_IMAGENES_VISION = 16   # si no marcas imágenes y la IA puede verlas, elige entre las primeras 16
# Respuesta acotada: la guía en JSON ocupa ~1.500 tokens; el resto queda para el razonamiento del modelo.
# El material que recibe cada proveedor depende de su límite (ia.PROVEEDORES).
MAX_TOKENS_RESPUESTA = 5000

PROMPT_PIE = """Eres el ayudante de {docente} de {asignatura}{colegio} (Chile), para crear guías PIE.
Las guías son para estudiantes NO LECTORAS de {nivel}, algunas con Síndrome de Down. Todo debe ser muy simple, visual y con poco texto.

Responde SOLO con un objeto JSON válido con esta forma:
{
  "objetivo": "COMPRENDER ...",
  "secciones": [
    {
      "titulo": "SUBTÍTULO CORTO",
      "imagen": "img3" o null,
      "pie_imagen": "QUÉ MUESTRA LA IMAGEN",
      "foto_web": {"busqueda": "términos para Wikimedia Commons", "pie": "PIE DE FOTO"} o null,
      "escribir": {"instruccion": "ESCRIBIR LOS DATOS DE ...", "frases": ["FRASE CORTA.", "..."]},
      "marcar": {"instruccion": "MARCAR LA RESPUESTA CORRECTA.", "pregunta": "¿...?", "opciones": ["...", "...", "..."], "correcta": 0} o null,
      "completar": {"instruccion": "COMPLETAR EL ESPACIO EN BLANCO.", "frases": ["FRASE CORTA."]} o null
    }
  ],
  "reconocer_unir": {"titulo": "RECONOZCO ...", "instruccion_reconocer": "RECONOCER LAS IMÁGENES.", "instruccion_unir": "UNIR CADA IMAGEN CON SU NOMBRE.", "pares": [{"imagen": "img1", "concepto": "CONCEPTO"}]} o null,
  "para_pensar": {"titulo": "PARA PENSAR", "imagen": null, "foto_web": null, "escribir": {...}, "marcar": {...}, "pregunta_conversar": "¿...?"} o null,
  "conceptos_clave": ["CONCEPTO.", "..."],
  "cierre": {"titulo": "PARA TERMINAR", "tipo": "dibujar" o "escribir", "instruccion": "DIBUJAR ..." o "ESCRIBIR ...", "frases": ["..."]},
  "picto_extra": {"PALABRA": "término simple en español para buscar en ARASAAC"}
}

REGLAS:
- Todo el texto en MAYÚSCULAS, con ortografía y tildes correctas (las mayúsculas también llevan tilde: ESTÁ, GANÓ, ¿QUÉ?). Frases muy cortas (máximo 8 palabras), sujeto + verbo + complemento, vocabulario concreto.
- El objetivo parte con un verbo en infinitivo (COMPRENDER, CONOCER, RECONOCER, IDENTIFICAR...).
- De 4 a 7 secciones de contenido, cada una con un subtítulo corto. Cada sección: 2 a 5 frases para ESCRIBIR, un MARCAR y un COMPLETAR.
- Las instrucciones empiezan SIEMPRE con su verbo: ESCRIBIR, COMPLETAR, MARCAR, RECONOCER, UNIR, COPIAR (o DIBUJAR en el cierre).
- La estudiante nunca debe adivinar: la respuesta correcta de cada MARCAR debe aparecer escrita tal cual en una frase del ESCRIBIR de esa misma sección o verse en su imagen. Tres opciones cortas, una correcta ("correcta" es su posición, empezando en 0).
- MARCAR siempre lleva imagen de referencia: cada sección con MARCAR debe tener "imagen" o "foto_web".
- COMPLETAR repite una frase del ESCRIBIR de la misma sección.
- Imágenes: usa solo los identificadores de la lista de imágenes del material (img1, img2...), según su descripción, y no repitas una misma imagen en dos secciones. Si no hay imagen adecuada, usa "foto_web" con una búsqueda concreta para Wikimedia Commons (ej: "volcán Osorno", "célula animal microscopio", "mapa de Chile").
- RECONOCER/UNIR: 3 o 4 pares con imágenes del material ya usadas en las secciones y su concepto (2 a 4 palabras). Si hay menos de 2 imágenes del material, usa null.
- "para_pensar" es opcional: una reflexión simple sobre el tema, con una pregunta para conversar en clase.
- "conceptos_clave": 5 a 10 conceptos del tema, terminados en punto.
- "cierre": un recuadro para dibujar o un ESCRIBIR final, según lo que mejor cierre el tema.
- "picto_extra": TODAS las palabras con contenido que aparecen en la guía (sustantivos, verbos, adjetivos importantes, lugares y pueblos), escritas sin tildes y en mayúscula tal como aparecen, con el término simple que mejor las representa en ARASAAC: verbos en infinitivo ("PERDIO": "perder", "LLEGARON": "llegar"), sustantivos en singular ("TIERRAS": "tierra"), conceptos abstractos con algo concreto que los represente ("ECONOMIA": "dinero", "FRACCION": "pizza", "ENERGIA": "luz"). ARASAAC no tiene nombres propios: los lugares van con "mapa", "ciudad", "isla" o "montaña" ("VALDIVIA": "ciudad", "CHILOE": "isla") y los pueblos con "indígena" o "pueblo originario". Normalmente 30 a 60 palabras.
- Usa solo información del material entregado y de las indicaciones del docente."""


def prompt_pie(perfil, asignatura, nivel):
    """Prompt de la guía con los datos del perfil (docente y colegio) y la asignatura elegida."""
    docente = (perfil.tratamiento_docente.strip() or "el profesor") + (f" {perfil.nombre_docente}" if perfil.nombre_docente else "")
    valores = {"{docente}": docente, "{asignatura}": asignatura or "la asignatura",
               "{colegio}": f" del {perfil.colegio}" if perfil.colegio.strip() else "", "{nivel}": nivel}
    texto = PROMPT_PIE
    for marca, valor in valores.items():   # replace y no format: el prompt contiene llaves de JSON
        texto = texto.replace(marca, valor)
    return texto


# ---------- Extracción de texto e imágenes ----------

def _normalizar_imagen(pil):
    """Convierte a PNG/JPEG reducido; devuelve (bytes, ext, ancho, alto) o None si no sirve."""
    try:
        pil.load()
        ancho, alto = pil.size
        if min(ancho, alto) < LADO_MINIMO or max(ancho, alto) / min(ancho, alto) > 5:
            return None
        if max(ancho, alto) > LADO_MAXIMO:
            pil.thumbnail((LADO_MAXIMO, LADO_MAXIMO))
        salida = BytesIO()
        if pil.mode in ("RGBA", "LA", "P"):
            pil.convert("RGBA").save(salida, "PNG", optimize=True)
            ext = "png"
        else:
            pil.convert("RGB").save(salida, "JPEG", quality=85)
            ext = "jpg"
        return salida.getvalue(), ext, pil.size[0], pil.size[1]
    except Exception:
        return None


def _formas_con_imagen(formas):
    for forma in formas:
        if hasattr(forma, "shapes"):  # grupo
            yield from _formas_con_imagen(forma.shapes)
        else:
            try:
                yield forma, forma.image.blob
            except Exception:
                continue


def extraer_material(archivos, avisos=None):
    """archivos: [(nombre, bytes)]. Devuelve (texto, imagenes) de PDF, PowerPoint y Word; los avisos se agregan a la lista."""
    avisos = avisos if avisos is not None else []
    textos, crudas = [], []  # crudas: (PIL.Image, origen)
    for nombre, contenido in archivos:
        archivo = BytesIO(contenido)
        bajo = nombre.lower()
        try:
            if bajo.endswith(".pdf"):
                lector = PdfReader(archivo)
                for n, pagina in enumerate(lector.pages, start=1):
                    textos.append(f"--- {nombre} · Página {n} ---\n{pagina.extract_text() or ''}")
                    try:
                        for im in pagina.images:
                            try:
                                crudas.append((im.image, f"{nombre} · pág. {n}"))
                            except Exception:
                                continue
                    except Exception:
                        continue
            elif bajo.endswith(".pptx"):
                pres = Presentation(archivo)
                for n, lamina in enumerate(pres.slides, start=1):
                    partes = [f.text_frame.text for f in lamina.shapes if getattr(f, "has_text_frame", False) and f.text_frame.text.strip()]
                    textos.append(f"--- {nombre} · Diapositiva {n} ---\n" + "\n".join(partes))
                    for _, blob in _formas_con_imagen(lamina.shapes):
                        try:
                            crudas.append((Image.open(BytesIO(blob)), f"{nombre} · diap. {n}"))
                        except Exception:
                            continue
            elif bajo.endswith(".docx"):
                doc = Document(archivo)
                partes = [p.text for p in doc.paragraphs if p.text.strip()]
                for tabla in doc.tables:
                    for fila in tabla.rows:
                        partes.append(" | ".join(c.text.strip() for c in fila.cells if c.text.strip()))
                textos.append(f"--- {nombre} ---\n" + "\n".join(partes))
                for rel in doc.part.rels.values():
                    if "image" in rel.reltype:
                        try:
                            crudas.append((Image.open(BytesIO(rel.target_part.blob)), nombre))
                        except Exception:
                            continue
        except Exception as e:
            avisos.append(f"No se pudo leer {nombre}: {e}")

    imagenes, vistas = [], set()
    for pil, origen in crudas:
        resultado = _normalizar_imagen(pil)
        if not resultado:
            continue
        datos, ext, ancho, alto = resultado
        huella = hashlib.md5(datos).hexdigest()
        if huella in vistas:
            continue
        vistas.add(huella)
        imagenes.append({"id": f"img{len(imagenes) + 1}", "datos": datos, "ext": ext, "origen": origen})
        if len(imagenes) >= MAX_IMAGENES:
            avisos.append(f"Se muestran solo las primeras {MAX_IMAGENES} imágenes.")
            break
    return "\n\n".join(textos), imagenes


def slug(texto):
    texto = unicodedata.normalize("NFD", texto or "")
    texto = "".join(c for c in texto if not unicodedata.combining(c)).lower()
    return re.sub(r"[^a-z0-9]+", "_", texto).strip("_") or "guia"


def leer_json(texto):
    """Lee el JSON del modelo; si viene con comillas escapadas de más (\\"), lo repara."""
    texto = texto or ""
    m = re.search(r"\{.*\}", texto, re.DOTALL)
    bloque = m.group(0) if m else texto
    # Reparación: comillas escapadas de más y objetos envueltos como texto ("{...}")
    reparado = bloque.replace('\\"', '"')
    reparado = re.sub(r'([}\]])"(?=\s*[,\]}])', r'\1', reparado)
    reparado = re.sub(r'(?<=[\[,])(\s*)"(?=\s*[{\[])', r'\1', reparado)
    ultimo_error = None
    for candidato in (texto, bloque, reparado):
        try:
            return _desenvolver(json.loads(candidato))
        except Exception as e:
            ultimo_error = e
    raise ultimo_error


def _desenvolver(valor):
    """Convierte en objetos los que el modelo dejó escritos como texto JSON ("{...}")."""
    if isinstance(valor, str) and valor.strip()[:1] in "{[" and valor.strip():
        try:
            return _desenvolver(json.loads(valor))
        except Exception:
            return valor
    if isinstance(valor, dict):
        return {k: _desenvolver(v) for k, v in valor.items()}
    if isinstance(valor, list):
        return [_desenvolver(v) for v in valor]
    return valor


def compactar_material(texto):
    """Quita espacios y líneas vacías del material."""
    texto = re.sub(r"[ \t]+", " ", texto or "")
    return re.sub(r"\n\s*\n+", "\n", texto).strip()


def pedir_guia(construir_mensajes):
    """Pide la guía en modo JSON (con paso automático a otro proveedor); si el JSON viene mal formado lo repara
    o la pide de nuevo sin el modo JSON estricto. Devuelve (spec, proveedor_usado, avisos)."""
    texto, usado, _, avisos = ia.completar(construir_mensajes, max_tokens=MAX_TOKENS_RESPUESTA, temperature=0.2,
                                           json=True, razonamiento_bajo=True)
    try:
        return leer_json(texto), usado, avisos
    except Exception:
        pass
    reintento = lambda info: construir_mensajes(info) + [{"role": "user", "content":
        "Responde solo con el objeto JSON, sin texto adicional y sin escapar las comillas."}]
    texto, usado, _, avisos2 = ia.completar(reintento, max_tokens=MAX_TOKENS_RESPUESTA, temperature=0.1,
                                            razonamiento_bajo=True)
    avisos = avisos + avisos2
    try:
        return leer_json(texto), usado, avisos
    except Exception:
        pass
    # Último recurso: la IA corrige la sintaxis de su propia respuesta, sin cambiar el contenido
    reparar = lambda info: [
        {"role": "system", "content": "Corriges la sintaxis de JSON. Devuelve SOLO el mismo objeto JSON, válido: comillas "
                                      "dobles escapadas dentro de textos, comas y llaves en su lugar. No cambies el contenido."},
        {"role": "user", "content": texto[:20000]}]
    texto, usado, _, avisos3 = ia.completar(reparar, max_tokens=MAX_TOKENS_RESPUESTA, temperature=0, json=True,
                                            razonamiento_bajo=True)
    return leer_json(texto), usado, avisos + avisos3


PROMPT_ORTOGRAFIA = """Eres corrector ortográfico de español de Chile. Recibes una lista JSON con los textos de una guía escolar (en mayúsculas).
Corrige SOLO errores de ortografía: tildes, letras equivocadas y palabras mal escritas o inexistentes (ej: RICHESAS -> RIQUEZAS, OSEANO -> OCÉANO).
No cambies el sentido, no reescribas ni acortes frases, no cambies nombres propios bien escritos (Chiloé, Selknam, Rapa Nui).
Responde solo con JSON: {"correcciones": [{"original": "texto tal como venía", "corregido": "texto corregido"}]}, incluyendo únicamente los textos que cambian. Si no hay errores: {"correcciones": []}."""

_CLAVES_SIN_REVISAR = {"imagen", "busqueda", "correcta", "picto_extra"}


def _textos_de_spec(valor, clave=None):
    """Todos los textos visibles de la guía (no ids de imagen, búsquedas web ni pictogramas)."""
    if clave in _CLAVES_SIN_REVISAR:
        return
    if isinstance(valor, str):
        if re.search(r"[A-Za-zÁÉÍÓÚÑáéíóúñ]{2}", valor):
            yield valor
    elif isinstance(valor, dict):
        for k, v in valor.items():
            yield from _textos_de_spec(v, k)
    elif isinstance(valor, list):
        for v in valor:
            yield from _textos_de_spec(v, clave)


def _reemplazar_en_spec(valor, mapa, clave=None):
    if clave in _CLAVES_SIN_REVISAR:
        return valor
    if isinstance(valor, str):
        return mapa.get(valor, valor)
    if isinstance(valor, dict):
        return {k: _reemplazar_en_spec(v, mapa, k) for k, v in valor.items()}
    if isinstance(valor, list):
        return [_reemplazar_en_spec(v, mapa, clave) for v in valor]
    return valor


def revisar_ortografia(spec):
    """Segunda pasada breve de la IA solo para ortografía. Devuelve (spec_corregido, ['ANTES → DESPUÉS', ...])."""
    textos = list(dict.fromkeys(_textos_de_spec(spec)))
    if not textos:
        return spec, []
    mensajes = [{"role": "system", "content": PROMPT_ORTOGRAFIA},
                {"role": "user", "content": json.dumps(textos, ensure_ascii=False)}]
    texto, _, _, _ = ia.completar(lambda info: mensajes, max_tokens=3000, temperature=0, json=True, razonamiento_bajo=True)
    mapa, cambios = {}, []
    for c in (leer_json(texto).get("correcciones") or []):
        if not isinstance(c, dict):
            continue
        original, corregido = str(c.get("original", "")), str(c.get("corregido", "")).strip()
        # Solo correcciones de ortografía: el texto debe existir y seguir siendo casi igual
        if original in textos and corregido and corregido != original and \
                difflib.SequenceMatcher(None, original.upper(), corregido.upper()).ratio() >= 0.75:
            mapa[original] = corregido
            cambios.append(f"{original} → {corregido}")
    return _reemplazar_en_spec(spec, mapa), cambios


# ---------- Generación ----------

def generar(perfil, material, asignatura, nivel, tema, indicaciones, seleccion, descripciones):
    """Pide la guía a la IA y revisa la ortografía.
    material: {texto, imagenes}; seleccion: ids de imágenes elegidas; descripciones: {id: qué muestra}.
    Devuelve dict {spec, usado, avisos_ia, recorte, ortografia}."""
    texto_material = compactar_material(material["texto"])
    imgs = {im["id"]: im for im in material["imagenes"]}
    elegidas = [imgs[i] for i in seleccion if i in imgs]
    recorte = {}

    def construir_mensajes(info_proveedor):
        limite = info_proveedor["max_material"]
        recorte["limite"] = limite if len(texto_material) > limite else None
        ve = info_proveedor["vision"]
        lista = "\n".join(
            f"- {im['id']}: {descripciones.get(im['id']) or ('ver la imagen adjunta' if ve else 'sin descripción; no usar')} ({im['origen']})"
            for im in elegidas) or "No hay imágenes del material; usa foto_web."
        usuario = (
            f"Nivel: {nivel}\nTema: {tema}\n\n"
            f"INDICACIONES DEL DOCENTE (prioridad máxima):\n{indicaciones.strip() or 'Sin indicaciones adicionales.'}\n\n"
            f"IMÁGENES DEL MATERIAL DISPONIBLES:\n{lista}\n\n"
            + ("Las imágenes van adjuntas a continuación, cada una precedida de su identificador. Míralas para decidir "
               "cuáles sirven (descarta logos, adornos y capturas con mucho texto) y escribe 'pie_imagen' según lo que "
               "realmente muestran.\n\n" if ve and elegidas else "")
            + f"TEXTO DEL MATERIAL:\n{texto_material[:limite] or 'Sin texto; usa el tema.'}")
        contenido = usuario
        if ve and elegidas:
            import base64
            contenido = [{"type": "text", "text": usuario}]
            for im in elegidas:
                mime = "image/png" if im["ext"] == "png" else "image/jpeg"
                contenido.append({"type": "text", "text": f"Imagen {im['id']}:"})
                contenido.append({"type": "image_url", "image_url": {
                    "url": f"data:{mime};base64,{base64.b64encode(im['datos']).decode()}"}})
        return [{"role": "system", "content": prompt_pie(perfil, asignatura, nivel)}, {"role": "user", "content": contenido}]

    spec, usado, avisos_ia = pedir_guia(construir_mensajes)
    try:
        spec, cambios = revisar_ortografia(spec)
    except Exception as e:
        cambios = [f"(No se pudo revisar la ortografía: {str(e)[:120]})"]
    return {"spec": spec, "usado": usado, "avisos_ia": avisos_ia, "recorte": recorte.get("limite"), "ortografia": cambios}


def armar_html(perfil, spec, nivel, asignatura, tema, material, seleccion):
    """Arma el HTML de la guía. Devuelve (html, avisos)."""
    from nucleo.armar_guia_pie import armar_guia
    imagenes = {im["id"]: (im["datos"], im["ext"]) for im in (material or {}).get("imagenes", []) if im["id"] in seleccion}
    return armar_guia(spec, NIVELES[nivel], imagenes, titulo_doc=f"Guía PIE - {tema} {nivel}",
                      encabezado=perfil.encabezado(asignatura))
